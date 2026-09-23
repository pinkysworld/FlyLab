#!/usr/bin/env python3
"""Export the reviewed Markdown manuscript to editable Word, without rerunning science.

Requires pandoc and python-docx. Run with the document-authoring Python runtime.
The four embedded result figures are numbered in reading order; other F/T IDs
remain supplementary artifact identifiers. Numerical source files are read only.
"""
from __future__ import annotations

import re
import hashlib
import json
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
PAPERS = ROOT / "papers"
OUTPUT = PAPERS / "submission" / "FlyLab_manuscript.docx"
# Section -> archived figure ID. These are deliberately a small main-text set.
FIGURES = {"3.2": "F6", "3.3": "F13", "3.4": "F14", "3.6": "F11"}


def manuscript_markdown() -> str:
    source = (PAPERS / "IJRC_FlyLab_draft.md").read_text()
    source = source.split("### Figures and tables")[0].rstrip().removesuffix("---").rstrip()
    # Submission status and build explanations belong in the repository, not title matter.
    source = re.sub(r"^\*\*Computational-methods Research Article\..*?\n", "", source, flags=re.M)
    source = re.sub(r"^> Simulation-result numeric spans.*?\n", "> No number in this paper was measured in a living animal.\n", source, flags=re.M)
    source = source.replace("https://github.com/pinkysworld/FlyLab\nSoftware version", "https://github.com/pinkysworld/FlyLab. Software version")
    captions = dict(re.findall(r"^\*\*(F\d+_[^.]+)\.\*\* (.+)$", (PAPERS / "figures/captions.md").read_text(), re.M))
    sections = re.split(r"(?=^### 3\.)", source, flags=re.M)
    main_ids = {identifier: str(i + 1) for i, identifier in enumerate(FIGURES.values())}
    for i, section in enumerate(sections):
        match = re.match(r"### (3\.\d+)\b", section)
        if not match or match[1] not in FIGURES:
            continue
        identifier = FIGURES[match[1]]
        name = next(key for key in captions if key.split("_")[0] == identifier)
        figure = f"\n\n![{name}]({PAPERS / 'figures' / (name + '.png')})\n\n**Figure {main_ids[identifier]}.** {captions[name]}\n\n"
        # Last result section shares a chunk with Discussion; insert before it.
        boundary = re.search(r"^## ", section, re.M)
        if boundary:
            sections[i] = section[:boundary.start()].rstrip() + figure + section[boundary.start():]
        else:
            sections[i] = section.rstrip() + figure
    source = "".join(sections)
    # Only references in prose, never paths or alt text, get remapped.
    lines = []
    for line in source.splitlines():
        if not line.startswith("!["):
            line = re.sub(r"\bF(\d+)\b", lambda m: main_ids.get(m[0], "S" + m[0]), line)
            line = re.sub(r"\bT(\d+[a-z]?)\b", r"ST\1", line)
        lines.append(line)
    source = "\n".join(lines)
    source = source.replace("| Graph model |", "**Table II. Graph degradations and retained information.**\n\n| Graph model |", 1)
    # DOI strings become usable links without changing citation text or numbering.
    source = re.sub(r"DOI (10\.[^\s]+?)(?=\.?(?:\s|$))", lambda m: "DOI [" + m[1].rstrip(".") + "](https://doi.org/" + m[1].rstrip(".") + ")", source)
    note = ("Supplementary figures and tables retain the archive identifiers SF and ST. "
            "Their captions, data and the supplementary methods are available with the "
            "[supporting artifacts](https://github.com/pinkysworld/FlyLab/tree/codex/flylab-reviewed-manuscript/papers). "
            "Figures 1–4 correspond to archive figures F6, F13, F14 and F11, respectively.\n\n")
    source = source.replace("## References", note + "## References", 1)
    return source


def style_document(doc: Document) -> None:
    styles = {style.name: style for style in doc.styles}
    for section in doc.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Mm(20)
        section.left_margin = section.right_margin = Mm(20)
    for style in doc.styles:
        if style.type in (1, 2):
            style.font.name = "Times New Roman"
            style.font.color.rgb = RGBColor(0, 0, 0)
    normal = styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 18), ("Heading 1", 13), ("Heading 2", 12), ("Heading 3", 11)):
        style = styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    for name in ("Source Code", "Verbatim Char"):
        if name in doc.styles:
            doc.styles[name].font.name = "DejaVu Sans Mono"
            doc.styles[name].font.size = Pt(9)
    paragraphs = doc.paragraphs
    for i, paragraph in enumerate(paragraphs):
        if i == 0:
            paragraph.style = doc.styles["Title"]
        paragraph.paragraph_format.widow_control = True
        if paragraph._p.xpath(".//w:drawing"):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.alignment = 1
        if re.match(r"Figure \d+\.", paragraph.text):
            paragraph.paragraph_format.space_after = Pt(10)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
        if re.match(r"Table [IVX]+\.", paragraph.text):
            paragraph.paragraph_format.keep_with_next = True
        if re.match(r"\[\d+\]", paragraph.text):
            paragraph.paragraph_format.space_after = Pt(4)
            for run in paragraph.runs:
                run.font.size = Pt(10)
    for shape in doc.inline_shapes:
        ratio = min(Mm(170) / shape.width, Mm(100) / shape.height)
        shape.width = int(shape.width * ratio)
        shape.height = int(shape.height * ratio)
    for table in doc.tables:
        table.autofit = False
        n = len(table.columns)
        widths = [Mm(78), Mm(92)] if n == 2 else [Mm(50)] + [Mm(120 / (n - 1))] * (n - 1)
        for column, width in zip(table.columns, widths):
            column.width = width
        for r, row in enumerate(table.rows):
            if r == 0:
                marker = OxmlElement("w:tblHeader")
                row._tr.get_or_add_trPr().append(marker)
            for c, cell in enumerate(row.cells):
                cell.width = widths[c]
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for edge in ("top", "left", "bottom", "right"):
                    el = OxmlElement("w:" + edge)
                    for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                        el.set(qn("w:" + key), value)
                    borders.append(el)
                props.append(borders)
                margins = OxmlElement("w:tcMar")
                for edge in ("top", "left", "bottom", "right"):
                    el = OxmlElement("w:" + edge)
                    el.set(qn("w:w"), "80")
                    el.set(qn("w:type"), "dxa")
                    margins.append(el)
                props.append(margins)
                if r == 0:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "E7E6E6")
                    props.append(shade)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.space_before = Pt(3)
                    paragraph.paragraph_format.keep_with_next = False
                    for run in paragraph.runs:
                        run.font.size = Pt(9 if n > 3 else 10)
                        if r == 0:
                            run.font.bold = True
    doc.core_properties.author = "Michél Nguyen"
    doc.core_properties.last_modified_by = "Michél Nguyen"
    doc.core_properties.title = paragraphs[0].text
    doc.core_properties.comments = ""


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="flylab-word-") as temporary:
        md = Path(temporary) / "manuscript.md"
        md.write_text(manuscript_markdown())
        subprocess.run(["pandoc", str(md), "--from=markdown-implicit_figures+tex_math_dollars", "--to=docx", "--output", str(OUTPUT)], check=True)
    doc = Document(OUTPUT)
    style_document(doc)
    doc.save(OUTPUT)
    assert len(doc.inline_shapes) == 4
    assert not any("[[MISSING:" in p.text for p in doc.paragraphs)
    inputs = [PAPERS / "IJRC_FlyLab_draft.md", PAPERS / "figures/captions.md", Path(__file__)]
    inputs.extend(next((PAPERS / "figures").glob(identifier + "_*.png")) for identifier in FIGURES.values())
    record = {
        "operation": "Word export only; numerical analyses not rerun",
        "inputs_sha256": {str(path.relative_to(ROOT)): hashlib.sha256(path.read_bytes()).hexdigest() for path in inputs},
        "output_sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
        "main_figure_archive_ids": list(FIGURES.values()),
    }
    OUTPUT.with_suffix(".json").write_text(json.dumps(record, indent=2) + "\n")
    print(OUTPUT)


if __name__ == "__main__":
    main()
